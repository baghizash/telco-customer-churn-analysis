#!/usr/bin/env python3
"""Build Portofolio-Analisis-Telco-Churn.pptx (16:9, 10 slides, ID).
Redesign design-taste: Tema T2 (oat, blue, pink), varian Editorial.
Headline Liberation Serif bold di atas ground oat, blue untuk struktur,
pink untuk highlight churn, angka Noto Sans. Konten dan semua angka
identik dengan deck lama (build_ppt_legacy.py)."""
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

GROUND = RGBColor(0xEF, 0xE7, 0xD8)    # oat
SURFACE = RGBColor(0xF7, 0xF1, 0xE5)   # permukaan terang
INK = RGBColor(0x2E, 0x31, 0x38)
BLUE = RGBColor(0x4E, 0x7F, 0xA6)      # primary
BLUEDK = RGBColor(0x33, 0x53, 0x6E)
PINK = RGBColor(0xC7, 0x5B, 0x7F)      # aksen churn
PINKSOFT = RGBColor(0xEF, 0xC3, 0xD2)
MUTED = RGBColor(0x5F, 0x63, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SERIF = "Liberation Serif"
SANS = "Noto Sans"
CHARTS = "/home/hatch/workspace/telco-churn/charts"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def solid_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, x, y, w, h, color):
    sh = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    style_el = sh._element.find(qn("p:style"))
    if style_el is not None:
        sh._element.remove(style_el)
    return sh


def textbox(slide, x, y, w, h):
    tx = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tx.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf


def add_par(tf, text, size, color, bold=False, italic=False, first=False,
            align=None, font=SANS, spacing=1.0, space_after=8):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.line_spacing = spacing
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = color
    return p


def add_bullet(tf, text, size=16.5, first=False, color=INK, space_after=11):
    p = add_par(tf, text, size, color, first=first, space_after=space_after)
    p.line_spacing = 1.12
    pPr = p._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "\u2022"}))
    pPr.set("marL", "171450")
    pPr.set("indent", "-171450")
    return p


def bullets(slide, x, y, w, h, items, size=16.5):
    tf = textbox(slide, x, y, w, h)
    for i, b in enumerate(items):
        add_bullet(tf, b, size=size, first=(i == 0))
    return tf


def title(slide, text):
    rect(slide, 0.7, 0.55, 0.16, 0.52, PINK)
    tf = textbox(slide, 1.05, 0.38, 11.6, 0.95)
    add_par(tf, text, 32, INK, bold=True, first=True, font=SERIF, space_after=0)


def chart(slide, name, x, y, w):
    path = f"{CHARTS}/{name}.png"
    iw, ih = Image.open(path).size
    h = w * ih / iw
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return h


def caption(slide, text, x, y, w):
    tf = textbox(slide, x, y, w, 0.4)
    add_par(tf, text, 12, MUTED, italic=True, first=True,
            align=PP_ALIGN.CENTER, space_after=0)


def stat_block(slide, x, y, w, value, label, vcolor, vsize=30, lsize=13.5):
    tf = textbox(slide, x, y, w, 1.15)
    add_par(tf, value, vsize, vcolor, bold=True, first=True, space_after=2)
    add_par(tf, label, lsize, INK, space_after=0)


# ---------- Slide 1: Cover (Asymmetric Split + Spotlight Stat) ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
rect(s, 8.55, 0, 4.783, 7.5, BLUE)
rect(s, 8.43, 0, 0.12, 7.5, PINK)
tf = textbox(s, 0.85, 1.5, 7.35, 2.2)
add_par(tf, "Analisis Telco Customer Churn", 37, INK, bold=True, first=True,
        font=SERIF, spacing=1.04, space_after=0)
add_par(tf, "dengan SQL dan Python", 37, INK, bold=True,
        font=SERIF, spacing=1.04, space_after=0)
tf = textbox(s, 0.85, 3.8, 7.1, 1.4)
add_par(tf, "Studi Kasus: IBM Sample Telco Customer Churn. Siapa yang Pergi, "
            "Kenapa, dan Berapa Pendapatan yang Hilang",
        17, BLUEDK, first=True, spacing=1.25, space_after=0)
tf = textbox(s, 0.85, 6.05, 7.3, 1.0)
add_par(tf, "BAGHIZ ASSHIDDIQIE ARASYA", 17, INK, bold=True, first=True, space_after=3)
add_par(tf, "Data Analyst Portfolio Project", 13.5, MUTED, space_after=0)
tf = textbox(s, 9.0, 2.05, 3.95, 3.6)
add_par(tf, "26,6%", 80, WHITE, bold=True, first=True, space_after=6)
add_par(tf, "churn rate keseluruhan", 16, WHITE, bold=True, space_after=6)
add_par(tf, "1.869 dari 7.032 pelanggan berhenti berlangganan", 14.5, WHITE,
        spacing=1.2, space_after=0)

# ---------- Slide 2: Latar Belakang ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
title(s, "Latar Belakang")
bullets(s, 0.7, 1.6, 7.5, 5.2, [
    "Dataset IBM Sample Telco Customer Churn: 7.032 pelanggan bersih dengan 21 kolom demografi, layanan, kontrak, dan tagihan",
    "Pertanyaan utama: pelanggan mana yang berhenti berlangganan, kenapa mereka pergi, dan berapa pendapatan yang hilang",
    "Churn rate keseluruhan 26,6 persen atau 1.869 pelanggan, lebih dari 1 dari 4 pelanggan pergi",
    "Tujuan analisis: menemukan segmen paling rawan churn lalu menyusun rekomendasi retensi yang konkret dan terukur",
])
rect(s, 8.6, 1.6, 4.03, 4.95, SURFACE)
stat_block(s, 8.95, 1.95, 3.4, "7.032", "pelanggan bersih dianalisis", BLUE)
stat_block(s, 8.95, 3.3, 3.4, "26,6%", "churn rate keseluruhan", PINK)
stat_block(s, 8.95, 4.65, 3.4, "1.869", "pelanggan berhenti berlangganan", BLUE)

# ---------- Slide 3: Data dan Metodologi ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
title(s, "Data dan Metodologi")
bullets(s, 0.7, 1.6, 6.1, 5.2, [
    "Data dimuat ke SQLite lalu dianalisis dengan query SQL asli: agregasi, grouping, dan perbandingan antar segmen",
    "Pandas dipakai untuk eksplorasi lanjutan dan pembuatan 12 chart visual",
    "Churn didefinisikan sebagai kolom Churn bernilai Yes; rate dihitung sebagai Yes dibagi total pelanggan",
    "Segmen yang dianalisis: kontrak, masa berlangganan, layanan internet, metode pembayaran, tagihan, layanan dukungan, demografi, dan segmen kombinasi",
], size=16)
rect(s, 7.2, 1.6, 5.43, 4.7, PINKSOFT)
tf = textbox(s, 7.6, 1.95, 4.65, 4.1)
add_par(tf, "11", 54, PINK, bold=True, first=True, space_after=2)
add_par(tf, "baris dihapus dari data, bukan diimputasi", 16, INK, bold=True,
        spacing=1.15, space_after=8)
add_par(tf, "Pembersihan penting: kolom TotalCharges tiba sebagai teks dengan "
            "11 nilai kosong tersembunyi (semua tenure 0, belum pernah ditagih), "
            "11 baris itu dihapus karena imputasi berarti mengarang data",
        15, INK, spacing=1.25, space_after=0)

# ---------- Slide 4: Temuan 1, kontrak ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
title(s, "Temuan 1: Kontrak Bulanan Paling Berisiko")
bullets(s, 0.7, 1.6, 5.5, 3.6, [
    "Kontrak bulanan (month-to-month) churn 42,7 persen, 15 kali lipat kontrak 2 tahun yang hanya 2,8 persen",
    "Kontrak bulanan mencakup 55,1 persen pelanggan tetapi menyumbang 88,6 persen dari semua churn",
    "Kontrak 1 tahun churn 11,3 persen, jauh lebih aman dibanding bulanan",
    "Tanpa komitmen jangka panjang, pelanggan pergi kapan saja dan hampir tanpa hambatan",
], size=16)
stat_block(s, 0.7, 5.6, 2.6, "42,7%", "kontrak bulanan", PINK, vsize=32, lsize=13)
stat_block(s, 3.5, 5.6, 2.6, "2,8%", "kontrak 2 tahun", BLUE, vsize=32, lsize=13)
h = chart(s, "02_churn_contract", 6.7, 1.6, 5.93)
caption(s, "Churn rate per jenis kontrak", 6.7, 1.6 + h + 0.12, 5.93)

# ---------- Slide 5: Temuan 2, pelanggan baru ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
title(s, "Temuan 2: Tahun Pertama Paling Kritis")
h = chart(s, "03_churn_tenure", 0.7, 1.85, 5.93)
caption(s, "Churn rate per masa berlangganan", 0.7, 1.85 + h + 0.12, 5.93)
tf = textbox(s, 7.05, 1.6, 5.58, 1.3)
add_par(tf, "47,7%", 44, PINK, bold=True, first=True, space_after=2)
add_par(tf, "pelanggan baru pergi di tahun pertama", 14.5, INK, space_after=0)
bullets(s, 7.05, 3.05, 5.58, 3.9, [
    "Pelanggan baru 0 sampai 12 bulan churn 47,7 persen, hampir separuh pergi di tahun pertama",
    "Pelanggan baru menyumbang 55,5 persen dari semua churn, padahal hanya 30,9 persen dari basis pelanggan",
    "Churn turun bertahap sampai 9,5 persen setelah 4 tahun berlangganan",
    "Siapa yang bertahan melewati tahun pertama cenderung setia, jadi onboarding awal adalah investasi terbaik",
], size=16)

# ---------- Slide 6: Temuan 3, fiber + pembayaran ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
title(s, "Temuan 3: Fiber Mahal tapi Ditinggalkan")
bullets(s, 0.7, 1.5, 5.85, 1.75, [
    "Pelanggan Fiber optic churn 41,9 persen, lebih dari dua kali DSL yang 19,0 persen",
    "Tagihan bulanan 71 sampai 90 USD churn 37,8 persen, tertinggi di antara band tagihan",
], size=15.5)
bullets(s, 6.78, 1.5, 5.85, 1.75, [
    "Pelanggan yang membayar dengan cek elektronik churn 45,3 persen, hampir 3 kali lipat kartu kredit otomatis 15,3 persen",
    "Sinyal masalah nilai: pelanggan membayar paling mahal tetapi merasa layanannya tidak sepadan, dan pembayaran manual membuat pergi terasa mudah",
], size=15.5)
h = chart(s, "04_churn_internet", 0.7, 3.3, 6.3)
caption(s, "Churn rate per layanan internet", 0.7, 3.3 + h + 0.1, 6.3)
stat_block(s, 7.6, 3.75, 4.9, "41,9%", "churn pelanggan Fiber optic", PINK, vsize=36)
stat_block(s, 7.6, 5.35, 4.9, "45,3%", "churn pembayar cek elektronik", PINK, vsize=36)

# ---------- Slide 7: Temuan 4, dukungan + demografi ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
title(s, "Temuan 4: Dukungan Melindungi, Lansia Rentan")
bullets(s, 0.7, 1.5, 5.85, 1.8, [
    "Pelanggan internet tanpa Tech Support churn 41,6 persen, dibanding yang memakai hanya 15,2 persen",
    "Tanpa Online Security churn 41,8 persen, dibanding yang memakai 14,6 persen",
], size=15.5)
bullets(s, 6.78, 1.5, 5.85, 1.8, [
    "Senior citizen churn 41,7 persen, pelanggan tanpa pasangan 33,0 persen, tanpa tanggungan 31,3 persen",
    "Gender hampir tidak berpengaruh: perempuan 27,0 persen vs laki-laki 26,2 persen",
], size=15.5)
h = chart(s, "07_churn_support", 1.85, 3.35, 9.63)
caption(s, "Churn: Tech Support dan Online Security", 1.85, 3.35 + h + 0.1, 9.63)

# ---------- Slide 8: Temuan 5, kombinasi + pendapatan ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
title(s, "Temuan 5: Risiko Menumpuk dan Pendapatan Hilang")
bullets(s, 0.7, 1.6, 5.5, 3.5, [
    "Fiber optic plus kontrak bulanan churn 54,6 persen, lebih dari separuh pergi",
    "Kontrak bulanan plus pelanggan baru churn 51,4 persen, kombinasi dua faktor rawan",
    "Pelanggan yang churn membayar rata-rata 74,44 USD per bulan vs 61,31 USD yang bertahan, yang pergi adalah pelanggan bernilai lebih tinggi",
], size=16)
h = chart(s, "10_churn_combo", 6.7, 1.6, 5.93)
caption(s, "Churn pada segmen kombinasi paling rawan", 6.7, 1.6 + h + 0.12, 5.93)
rect(s, 0.7, 5.7, 11.93, 1.15, BLUE)
tf = textbox(s, 1.05, 5.92, 3.6, 0.8)
add_par(tf, "139.131 USD", 30, WHITE, bold=True, first=True, space_after=0)
tf = textbox(s, 4.75, 5.9, 7.6, 0.85)
add_par(tf, "Pendapatan yang hilang akibat churn per bulan, atau 30,5 persen "
            "dari seluruh pendapatan bulanan", 15.5, WHITE, first=True,
        spacing=1.2, space_after=0)

# ---------- Slide 9: Rekomendasi ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
title(s, "Rekomendasi Retensi")
recs = [
    "Bangun program onboarding 12 bulan pertama: sapa proaktif di bulan ke-1, 3, dan 6 plus cek kepuasan dini",
    "Konversi kontrak bulanan ke kontrak tahunan dengan insentif: diskon loyalitas, harga terkunci, dan bonus layanan",
    "Audit nilai layanan Fiber optic: tinjau kualitas jaringan per area, harga, dan komunikasi manfaatnya",
    "Dorong auto-pay dan migrasi dari cek elektronik dengan insentif kecil",
    "Tawarkan Tech Support dan Online Security proaktif, bundel dukungan gratis 3 bulan untuk segmen rawan",
    "Perhatian khusus untuk lansia dan pelanggan yang hidup sendiri dengan jalur dukungan yang lebih personal",
]
for i, text in enumerate(recs):
    col, row = i % 2, i // 2
    x, y = 0.7 + col * 6.11, 1.55 + row * 1.72
    rect(s, x, y, 5.82, 1.54, SURFACE)
    rect(s, x + 0.22, y + 0.24, 0.14, 0.14, PINK)
    tf = textbox(s, x + 0.55, y + 0.2, 5.05, 1.2)
    add_par(tf, text, 14.5, INK, first=True, spacing=1.18, space_after=0)

# ---------- Slide 10: Penutup (Designed Close) ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
title(s, "Kesimpulan")
bullets(s, 0.7, 1.6, 7.5, 4.6, [
    "Kontrak bulanan, tahun pertama, dan fiber optic adalah tiga pendorong churn terbesar",
    "Fokus retensi pada segmen kombinasi: fiber plus kontrak bulanan (54,6 persen) dan kontrak bulanan plus pelanggan baru (51,4 persen)",
    "Churn bukan sekadar angka pelanggan: 139.131 USD pendapatan hilang setiap bulan",
    "Tech stack: SQL di SQLite, Python dengan pandas, matplotlib, dan seaborn",
], size=16)
rect(s, 8.6, 1.6, 4.03, 4.6, SURFACE)
tf = textbox(s, 8.95, 1.95, 3.35, 3.9)
add_par(tf, "26,6%", 40, PINK, bold=True, first=True, space_after=4)
add_par(tf, "satu dari empat pelanggan pergi. Retensi segmen rawan adalah "
            "tuas pendapatan terbesar.", 14.5, INK, spacing=1.25, space_after=0)
tf = textbox(s, 0.7, 6.35, 8.6, 0.95)
add_par(tf, "BAGHIZ ASSHIDDIQIE ARASYA", 22, INK, bold=True, first=True,
        font=SERIF, space_after=3)
add_par(tf, "github.com/baghizash/telco-customer-churn-analysis", 14.5, BLUEDK,
        space_after=0)
tf = textbox(s, 8.6, 6.55, 4.03, 0.6)
add_par(tf, "Mari diskusi soal data dan retensi pelanggan.", 14.5, PINK,
        bold=True, first=True, space_after=0)

prs.save("/home/hatch/workspace/your_files/Portofolio-Analisis-Telco-Churn.pptx")
prs.save("/home/hatch/workspace/telco-churn/Portofolio-Analisis-Telco-Churn.pptx")
print("saved")
